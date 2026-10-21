"""The SpecSheet spoke: weight log and CG sheet -> the ADH mass model. Every number is synthetic."""
from types import SimpleNamespace

import openpyxl
import pytest
import yaml

from aeroplan import adh_doc
from aeroplan.consumers import missing_for_flight
from aeroplan.datum import datum_checks
from aeroplan.spokes import Context, load
from test_openvsp import make_doc

sheet = load("SpecSheet")
OVERRIDE = "Subassembly weight override"

SETTINGS = {
    "workbook": "sheet.xlsx", "weights_sheet": "W", "positions_sheet": "P", "configuration": "test", "datum_x_m": 0.0,
    "strips_per_side": 4,
    "subassemblies": {
        "Body": {"node": "1.2.2.2", "instance": "fuselage", "cg_group": "body", "exclude_rows": ["tanks"],
                 "components": {"Part A": "a"}},
        "Wing": {"node": "1.2.2.3", "instance": "wing_main", "cg_group": "wing", "distribute": "wing_main"},
    },
}


def q(value, units):
    return {"value": value, "units": units}


def make_full_doc():
    """The test aircraft of the OpenVSP tests with a valid 1.2 reference (MRP at x = 0.2 m)."""
    doc = make_doc()
    adh_doc.find(doc, "wing_main")["muas"]["planform"]["position"]["xyz"][0] = 0.0     # ADR 0023: wing root LE at x = 0
    air = adh_doc.node(doc, "1.2")
    air["muas"]["reference"] = {
        "s_ref": q(0.9, "m^2"), "b_ref": q(2.8, "m"), "c_ref": q(0.33, "m"), "nominal_cg": {"xyz": [0.2, 0.0, 0.0]},
        "mrp": {"config_id": "test.nominal", "point": {"frame": "design", "xyz": [0.2, 0.0, 0.0], "units": "m"}}}
    return doc


def write_book(path, total_g=1000.0):
    wb = openpyxl.Workbook()
    w = wb.active
    w.title = "W"
    rows = [("Body", "Part A", 100, 1, 100, "Weighed the physical item"),
            (None, "Part B", 50, 1, 50, "I made it up"),
            (None, "Part C", None, 1, 0, "Not Weighed"),
            (None, "Component Sum Weight", 150, 1, 150, None),
            (None, "Override Weight", 400, 1, 400, OVERRIDE),
            ("Wing", "Spar", 30, 2, 60, "Determined from CAD"),
            (None, "Unitemised Weight", 540, 1, 540, None),
            (None, "Override Weight", 600, 1, 600, OVERRIDE),
            ("Total", None, total_g, None, None, None)]
    for i, r in enumerate(rows, 9):
        for c, v in enumerate(r, 1):
            w.cell(i, c).value = v
    p = wb.create_sheet("P")
    for i, r in enumerate([("body", "a", None, None, 100, -200), (None, "b", None, None, 50, 100),
                           (None, "tanks", None, None, 500, 0), ("wing", "skin", None, None, 10, -50)], 3):
        for c, v in enumerate(r, 1):
            p.cell(i, c).value = v
    wb.save(path)


def test_classify():
    assert sheet.classify("Weighed from unfinished physical item") == {"source": "weighed", "placeholder": False}
    assert sheet.classify("Determined from CAD") == {"source": "cad", "placeholder": False}
    assert sheet.classify("Calculated from physical item (x)") == {"source": "estimate", "placeholder": False}
    for t in ("I made it up", "Not Weighed", None):
        assert sheet.classify(t) == {"source": "estimate", "placeholder": True}


def test_read_workbook(tmp_path):
    write_book(tmp_path / "sheet.xlsx")
    d = sheet.read_workbook(tmp_path / "sheet.xlsx", SETTINGS)
    body = d["groups"]["Body"]
    assert body["total_g"] == 400.0 and body["total_row"] == 13
    assert [p["name"] for p in body["parts"]] == ["Part A", "Part B"]
    assert body["unweighed"] == ["Part C"]
    assert [p["placeholder"] for p in body["parts"]] == [False, True]
    assert d["groups"]["Wing"]["parts"][0]["source"] == "cad"
    assert d["positions"]["body"]["a"]["x_mm"] == -200.0


def test_the_sheet_total_must_match(tmp_path):
    write_book(tmp_path / "sheet.xlsx", total_g=1000.5)
    with pytest.raises(ValueError, match="Total row"):
        sheet.read_workbook(tmp_path / "sheet.xlsx", SETTINGS)


def test_a_group_without_a_total_is_refused(tmp_path):
    write_book(tmp_path / "sheet.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "sheet.xlsx")
    wb["W"].cell(13, 6).value = None
    wb.save(tmp_path / "sheet.xlsx")
    with pytest.raises(ValueError, match="no row with"):
        sheet.read_workbook(tmp_path / "sheet.xlsx", SETTINGS)


def planform_of(doc):
    return lambda inst: adh_doc.find(doc, inst)["muas"]["planform"]


def test_point_masses(tmp_path):
    write_book(tmp_path / "sheet.xlsx")
    doc = make_full_doc()
    d = sheet.read_workbook(tmp_path / "sheet.xlsx", SETTINGS)
    subs = {s["group"]: s["point_masses"] for s in sheet.build_point_masses(d, SETTINGS, planform_of(doc))}
    a, rest = subs["Body"]
    assert a["tag"] == "BODY-A" and a["mass_kg"] == pytest.approx(0.1) and a["xyz"][0] == pytest.approx(-0.2)
    assert a["placeholder"] is False
    # the rest: the weighed total less the mapped part, at the free row (b), not at the excluded tanks
    assert rest["mass_kg"] == pytest.approx(0.3) and rest["xyz"][0] == pytest.approx(0.1)
    assert rest["placeholder"] is True and "Part B" in rest["source"] and "Part C" in rest["source"]
    wing = subs["Wing"]
    assert len(wing) == 8
    assert sum(p["mass_kg"] for p in wing) == pytest.approx(0.6)
    assert sum(p["mass_kg"] * p["xyz"][0] for p in wing) / 0.6 == pytest.approx(-0.05)


def test_mapped_parts_over_the_total_are_refused(tmp_path):
    write_book(tmp_path / "sheet.xlsx")
    d = sheet.read_workbook(tmp_path / "sheet.xlsx", SETTINGS)
    d["groups"]["Body"]["total_g"] = 90.0
    with pytest.raises(ValueError, match="weigh more"):
        sheet.build_point_masses(d, SETTINGS, planform_of(make_full_doc()))


def ingest(tmp_path):
    write_book(tmp_path / "sheet.xlsx")
    (tmp_path / "tools" / "specsheet" / "settings").mkdir(parents=True)
    (tmp_path / "tools" / "specsheet" / "settings" / "test.yaml").write_text(yaml.safe_dump(SETTINGS), encoding="utf-8")
    doc = make_full_doc()
    ctx = Context("test", tmp_path, "SpecSheet", SimpleNamespace(doc=doc))
    return doc, sheet.ingest(doc, ctx)


def test_ingest_writes_the_mass_model(tmp_path):
    doc, changes = ingest(tmp_path)
    assert changes
    air = adh_doc.node(doc, "1.2")["muas"]
    mp = air["mass_properties"]
    assert mp["mass"]["value"] == pytest.approx(1.0)
    assert mp["cg"]["xyz"][0] == pytest.approx(-0.02)      # (0.1*(-0.2) + 0.3*0.1 + 0.6*(-0.05)) / 1.0 (FRD: x = +F / 1000)
    assert mp["inertia"]["iyy"] > 0 and mp["placeholder"] is True
    assert len(air["point_masses"]) == 2 + 8
    assert sum(p["mass"]["value"] for p in air["point_masses"]) == pytest.approx(1.0)
    assert any(r["tool"] == "SpecSheet" for r in air["sources"])
    body = adh_doc.node(doc, "1.2.2.2")["muas"]["mass_properties"]
    assert body["mass"]["value"] == pytest.approx(0.4)
    # the reference (the MRP of the aero model) is not changed
    assert air["reference"]["mrp"]["point"]["xyz"] == [0.2, 0.0, 0.0]
    adh_doc.validate_document(doc)


def test_the_mrp_gap_is_a_warning_and_stops_aerolink(tmp_path):
    doc, _ = ingest(tmp_path)
    warns = [c for c in datum_checks(doc, tmp_path).checks if c.status == "warn" and "MRP differs" in c.name]
    assert len(warns) == 1
    assert any("muas.reference mrp" in m for m in missing_for_flight(doc))
    air = adh_doc.node(doc, "1.2")["muas"]
    air["reference"]["mrp"]["point"]["xyz"] = list(air["mass_properties"]["cg"]["xyz"])
    air["reference"]["nominal_cg"] = {"xyz": list(air["mass_properties"]["cg"]["xyz"])}
    assert not any("muas.reference mrp" in m for m in missing_for_flight(doc))


def test_a_datum_other_than_zero_is_refused(tmp_path):
    write_book(tmp_path / "sheet.xlsx")
    d = sheet.read_workbook(tmp_path / "sheet.xlsx", SETTINGS)
    with pytest.raises(ValueError, match="ADR 0023"):
        sheet.build_point_masses(d, {**SETTINGS, "datum_x_m": -0.08}, planform_of(make_full_doc()))


def test_a_wing_not_at_x_zero_is_refused(tmp_path):
    write_book(tmp_path / "sheet.xlsx")
    (tmp_path / "tools" / "specsheet" / "settings").mkdir(parents=True)
    (tmp_path / "tools" / "specsheet" / "settings" / "test.yaml").write_text(yaml.safe_dump(SETTINGS), encoding="utf-8")
    doc = make_full_doc()
    adh_doc.find(doc, "wing_main")["muas"]["planform"]["position"]["xyz"][0] = 0.1
    assert sheet._root_le_x(doc, SETTINGS) == pytest.approx(0.1)
    with pytest.raises(ValueError, match="ADR 0023"):
        sheet.ingest(doc, Context("test", tmp_path, "SpecSheet", SimpleNamespace(doc=doc)))


def test_a_matching_datum_gives_no_warning(tmp_path):
    doc, changes = ingest(tmp_path)
    assert sheet._root_le_x(doc, SETTINGS) == pytest.approx(0.0)
    assert not any(c.startswith("WARNING") for c in changes)


# ---------------------------------------------------------------------------------------------- battery trim
def trim_subs(tmp_path, **over):
    write_book(tmp_path / "sheet.xlsx")
    st = {**SETTINGS, "battery_trim": {"move": "BODY-A", "travel_m": 0.05, "neutral_point": "np.json", "source": "synthetic"}}
    st["battery_trim"].update(over)
    d = sheet.read_workbook(tmp_path / "sheet.xlsx", st)
    return st, sheet.build_point_masses(d, st, planform_of(make_full_doc()))


def band(x_np, lo=0.05, hi=0.15):
    return {"x_np": x_np, "c_ref": 0.33, "sm_min": lo, "sm_max": hi, "np_file": "np.json"}


def cg_of(subs):
    pms = [p for s in subs for p in s["point_masses"]]
    return sum(p["mass_kg"] * p["xyz"][0] for p in pms) / sum(p["mass_kg"] for p in pms)


def test_trim_inside_the_travel_gives_the_goal_margin(tmp_path):
    st, subs = trim_subs(tmp_path)                      # the synthetic CG is at x = -0.02 m (FRD), "BODY-A" weighs 0.1 kg
    rep = sheet.apply_battery_trim(subs, st, band(-0.02 - 0.10 * 0.33 - 0.003))        # the neutral point is behind the CG
    assert not rep["clamped"] and abs(rep["dx_m"]) < 0.05
    assert rep["sm"] == pytest.approx(0.10, abs=1e-12) and cg_of(subs) == pytest.approx(rep["cg_m"])
    assert "Battery cage trim" in subs[0]["point_masses"][0]["source"]


def test_trim_outside_the_travel_stops_at_the_limit(tmp_path):
    st, subs = trim_subs(tmp_path)
    x0 = subs[0]["point_masses"][0]["xyz"][0]
    rep = sheet.apply_battery_trim(subs, st, band(-0.02 - 0.25 * 0.33 - 0.02, 0.0, 0.5))      # goal 0.25, needs more than 50 mm
    assert rep["clamped"] and rep["dx_m"] == pytest.approx(-0.05)
    assert subs[0]["point_masses"][0]["xyz"][0] == pytest.approx(x0 - 0.05)
    assert rep["sm"] != pytest.approx(rep["sm_goal"]) and 0.0 <= rep["sm"] <= 0.5          # the margin reached is reported
    lines = sheet.trim_lines(rep, st)
    assert any(line.startswith("WARNING") for line in lines)
    st, subs = trim_subs(tmp_path)
    rep = sheet.apply_battery_trim(subs, st, band(-0.02 - 0.25 * 0.33 + 0.02, 0.0, 0.5))
    assert rep["dx_m"] == pytest.approx(0.05)


def test_trim_with_a_margin_outside_the_band_fails(tmp_path):
    st, subs = trim_subs(tmp_path)
    with pytest.raises(ValueError, match="outside the band"):
        sheet.apply_battery_trim(subs, st, band(-0.02 - 0.10 * 0.33 - 0.5))


def test_trim_needs_the_tag_and_the_file(tmp_path):
    st, subs = trim_subs(tmp_path, move="NOPE")
    with pytest.raises(ValueError, match="not in the mass model"):
        sheet.apply_battery_trim(subs, st, band(0.1))
    doc = make_full_doc()
    with pytest.raises(ValueError, match="neutral point file"):
        sheet.trim_inputs(doc, st, tmp_path)
    assert sheet.trim_inputs(doc, SETTINGS, tmp_path) is None            # no key: the old behaviour


# ------------------------------------------------------------------------------------ origin check (ADR 0023)
def origin_rows(doc, tmp_path):
    return [c for c in datum_checks(doc, tmp_path).checks if c.name.startswith("origin:")]


def test_the_origin_check_passes_at_x_zero(tmp_path):
    rows = origin_rows(make_full_doc(), tmp_path)
    assert len(rows) == 1 and rows[0].status == "pass"


def test_a_wing_at_x_001_fails_the_origin_check(tmp_path):
    doc = make_full_doc()
    adh_doc.find(doc, "wing_main")["muas"]["planform"]["position"]["xyz"][0] = 0.01
    rows = origin_rows(doc, tmp_path)
    assert rows[0].status == "fail" and not datum_checks(doc, tmp_path).ok


def test_a_tree_with_no_main_wing_skips_the_origin_check(tmp_path):
    doc = make_full_doc()
    adh_doc.find(doc, "wing_main")["muas"]["instance"] = "not_the_main_wing"      # stands for a multirotor or a reference model
    rows = origin_rows(doc, tmp_path)
    assert len(rows) == 1 and rows[0].status == "info" and "skipped" in rows[0].name
