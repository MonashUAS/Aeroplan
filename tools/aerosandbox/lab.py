"""Notebook helpers of the AeroSandbox spoke. They run in tools/aerosandbox/.venv (aerosandbox and aeroplan are both there).

    import sys; sys.path.insert(0, "..")      # the notebooks are in tools/aerosandbox/notebooks
    import lab
    L = lab.load("currawong")

The notebooks only read the ADH source tree. They write plots and tables to work/aerosandbox/<aircraft>/lab/ (git
ignores it). A change of the aircraft comes through a spoke.

`load()` gives a `Lab`:
    doc, acft     the ADH document and the source tree of the aircraft
    st            the spoke settings (tools/aerosandbox/settings/<aircraft>.yaml)
    spec          the tool-neutral spec of the spoke (spec.py)
    airplane      the asb.Airplane of the worker (asb_worker.build_airplane), reference point = the MRP
    masses        the point masses of node 1.2 [{tag, mass_kg, xyz, placeholder, source}]
    flow5         the DAVE-ML aero model of node 1.2 (aeroplan.daveml_eval.Model), or None
    configs       named mass configurations {name: point masses}: the nominal one and the tank layouts of the sheet
"""
from __future__ import annotations

import importlib.util
import math
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional

import numpy as np

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
G0 = 9.80665


def _worker():
    """asb_worker.py as a module (it is a script of the folder, not a package)."""
    name = "asb_worker_lab"
    if name not in sys.modules:
        spec = importlib.util.spec_from_file_location(name, HERE / "asb_worker.py")
        mod = importlib.util.module_from_spec(spec)
        sys.modules[name] = mod
        spec.loader.exec_module(mod)
    return sys.modules[name]


@dataclass
class Lab:
    aircraft: str
    acft: Any = field(repr=False)
    st: Dict[str, Any] = field(repr=False)
    spec: Dict[str, Any] = field(repr=False)
    airplane: Any
    masses: List[Dict[str, Any]] = field(repr=False)
    flow5: Any
    configs: Dict[str, List[Dict[str, Any]]] = field(repr=False)
    out: Path
    flow5_data: Optional[Dict[str, Any]] = field(default=None, repr=False)

    @property
    def doc(self) -> Dict[str, Any]:
        return self.acft.doc

    @property
    def ref(self) -> Dict[str, Any]:
        return self.spec["reference"]

    def plane(self, xyz_ref=None):
        """The airplane with another moment reference point (a copy)."""
        import copy

        new = copy.copy(self.airplane)
        new.xyz_ref = list(self.ref["xyz_ref"] if xyz_ref is None else xyz_ref)
        return new

    def flow5_const(self) -> Dict[str, float]:
        """The coefficients of the flow5 aero build-up (CLa, Cma, CLde, Cmde, CD0, K, CL0, Cm0, ... per rad)."""
        return self.flow5_eval(0.0)

    def flow5_eval(self, alpha_deg: float, **inputs: float) -> Dict[str, float]:
        """Every output of the flow5 aero model (CL, CD, Cm, ...). Inputs in the units of the DAVE-ML file (rad)."""
        base = {v["var_id"]: 0.0 for v in self.flow5.vars.values() if v.get("is_input")}
        return self.flow5.evaluate({**base, "alpha": math.radians(alpha_deg), **inputs})


def point_masses_of(doc: Dict[str, Any]) -> List[Dict[str, Any]]:
    from aeroplan import adh_doc

    air = adh_doc.node(doc, "1.2")["muas"]
    return [{"tag": p["tag"], "mass_kg": p["mass"]["value"], "xyz": list(p["position"]["xyz"]),
             "placeholder": bool(p.get("placeholder")), "source": p["source"]} for p in air.get("point_masses") or []]


def configurations(root: Path, aircraft: str, nominal: List[Dict[str, Any]]) -> Dict[str, Any]:
    """The nominal mass model and the layouts of the sheet that add the 1.3 kg tank (rows `tanks` of the two CG
    sheets: mass from column E, x = -F / 1000). The second sheet is a longer nose: its tank x is taken as written."""
    from aeroplan.spokes import context, load

    spoke = load("SpecSheet")
    ctx = context("SpecSheet", aircraft)
    st = ctx.settings()
    if not st or not (root / st["workbook"]).is_file():
        return {"nominal": nominal}
    import openpyxl

    wb = openpyxl.load_workbook(root / st["workbook"], data_only=True)
    cfgs = {"nominal (no payload)": nominal}
    for label, sheet_name in (("nominal + 1 forward tank", "CG Calcs 1 forward tank"),
                              ("nominal + tanks of the 2-tank sheet", "CG Calcs 2 tanks")):
        if sheet_name not in wb.sheetnames:
            continue
        row = spoke.read_positions(wb[sheet_name]).get("fuse", {}).get("tanks")
        if row:
            cfgs[label] = nominal + [{"tag": "TANKS", "mass_kg": row["mass_g"] / 1000.0, "xyz": [-row["x_mm"] / 1000.0, 0, 0],
                                      "placeholder": True, "source": f"{sheet_name}, row {row['row']} tanks"}]
    return cfgs


def load(aircraft: str = "currawong") -> Lab:
    """Load the aircraft from its source tree (read only)."""
    import yaml
    from aeroplan import adh_doc, source_tree
    from aeroplan.daveml_eval import Model
    from aeroplan.spokes import context, load as load_spoke

    acft = source_tree.read(ROOT / "aircraft" / aircraft)
    ctx = context("AeroSandbox", aircraft, acft)
    st = yaml.safe_load(ctx.settings_file.read_text(encoding="utf-8"))
    spoke = load_spoke("AeroSandbox")
    spec = spoke.asb_spec.spec_from_doc(acft.doc, st, ctx)
    airplane = _worker().build_airplane(spec)
    airplane.xyz_ref = list(spec["reference"]["xyz_ref"])
    air = adh_doc.node(acft.doc, "1.2")["muas"]
    ref_id = next((d["ref"] for d in air.get("datasets", []) if d["role"] == "aero"), None)
    out = ROOT / "work" / "aerosandbox" / aircraft / "lab"
    out.mkdir(parents=True, exist_ok=True)
    masses = point_masses_of(acft.doc)
    data = ctx.dataset(ref_id) if ref_id else None
    return Lab(aircraft, acft, st, spec, airplane, masses, Model(data) if data else None,
               configurations(ROOT, aircraft, masses), out, data)


# ------------------------------------------------------------------------------------------------------ the mass
def rollup(masses: List[Dict[str, Any]]) -> Dict[str, Any]:
    from aeroplan import mass_model

    return mass_model.rollup(masses)


def asb_mass(masses: List[Dict[str, Any]]):
    """The asb.MassProperties of a list of point masses (their sum)."""
    import aerosandbox as asb

    total = asb.MassProperties(mass=0.0)
    for p in masses:
        total = total + asb.MassProperties(mass=p["mass_kg"], x_cg=p["xyz"][0], y_cg=p["xyz"][1], z_cg=p["xyz"][2])
    return total


def group_of(tag: str) -> str:
    """The subassembly of a point mass tag: WINGS-S3 -> WINGS."""
    return tag.split("-")[0]


# ---------------------------------------------------------------------------------------------------- the geometry
def planform_numbers(wing) -> Dict[str, float]:
    """Area, span, MAC and the MAC leading edge of an asb.Wing (one side of the dihedral is counted, the wing is flat enough)."""
    mac = wing.mean_aerodynamic_chord()
    le = wing.mean_aerodynamic_chord_xyz_le() if hasattr(wing, "mean_aerodynamic_chord_xyz_le") else None
    out = {"area_m2": float(wing.area()), "span_m": float(wing.span()), "mac_m": float(mac), "AR": float(wing.aspect_ratio())}
    if le is not None:
        out["mac_le_x_m"] = float(le[0])
    return out


def figure_3d(lab: Lab, masses: Optional[List[Dict[str, Any]]] = None, show_cg: bool = True):
    """Plotly 3D view: the airplane (AeroSandbox) with the point masses (area ~ mass), the CG and the MRP."""
    import plotly.graph_objects as go

    masses = lab.masses if masses is None else masses
    fig = lab.airplane.draw(backend="plotly", show=False)
    groups: Dict[str, List[Dict[str, Any]]] = {}
    for p in masses:
        groups.setdefault(group_of(p["tag"]), []).append(p)
    for g, ps in groups.items():
        fig.add_trace(go.Scatter3d(
            x=[p["xyz"][0] for p in ps], y=[p["xyz"][1] for p in ps], z=[p["xyz"][2] for p in ps], mode="markers", name=g,
            text=[f"{p['tag']}: {p['mass_kg'] * 1000:.0f} g" for p in ps], hoverinfo="text",
            marker=dict(size=[3 + 14 * math.sqrt(p["mass_kg"]) for p in ps], opacity=0.75)))
    if show_cg:
        cg = rollup(masses)["cg"]
        fig.add_trace(go.Scatter3d(x=[cg[0]], y=[cg[1]], z=[cg[2]], mode="markers", name="CG",
                                   marker=dict(size=9, symbol="diamond", color="red")))
        m = lab.ref["xyz_ref"]
        fig.add_trace(go.Scatter3d(x=[m[0]], y=[m[1]], z=[m[2]], mode="markers", name="MRP (aero reference)",
                                   marker=dict(size=9, symbol="cross", color="black")))
    fig.update_layout(scene=dict(aspectmode="data"), height=620, margin=dict(l=0, r=0, t=30, b=0))
    return fig


def atmosphere(altitude_m: float = 0.0):
    import aerosandbox as asb

    return asb.Atmosphere(altitude=altitude_m)


def num(x) -> float:
    """A float from an AeroSandbox result (they come as 0-d or 1-element arrays)."""
    return float(np.ravel(x)[0])


# ------------------------------------------------------------------------------------ flight condition and aero runs
def condition(V: float, weight_N: float, s_ref: float, c_ref: float, altitude_m: float = 0.0) -> Dict[str, float]:
    """Air state and the lift coefficient of level flight at speed V."""
    atm = atmosphere(altitude_m)
    rho, mu = float(atm.density()), float(atm.dynamic_viscosity())
    q = 0.5 * rho * V ** 2
    return {"V": V, "rho": rho, "mu": mu, "q": q, "Re_c": rho * V * c_ref / mu, "CL_req": weight_N / (q * s_ref),
            "weight_N": weight_N, "s_ref": s_ref, "c_ref": c_ref, "altitude_m": altitude_m}


def subset(lab: Lab, wings=(), fuselage: bool = False, xyz_ref=None, bodies=None):
    """A copy of the airplane with some of the wings (by name) and the bodies. fuselage=True keeps every body (the
    fuselage and the tail boom); bodies=[names] keeps only those."""
    import copy

    p = lab.plane(xyz_ref)
    p = copy.copy(p)
    p.wings = [w for w in lab.airplane.wings if w.name in wings]
    if bodies is not None:
        p.fuselages = [f for f in lab.airplane.fuselages if f.name in bodies]
    else:
        p.fuselages = list(lab.airplane.fuselages) if fuselage else []
    return p


def run_aero(plane, cond: Dict[str, float], alpha_deg: float, method: str = "buildup", de_deg: float = 0.0,
             **vlm) -> Dict[str, float]:
    """CL, CD, Cm (about plane.xyz_ref) at one angle of attack. method: "buildup" (AeroBuildup), "vlm" (vortex
    lattice), "ll" (LiftingLine), "nll" (NonlinearLiftingLine). The keyword arguments go to the VLM and the lifting
    lines (resolution, align_trailing_vortices_with_wind). Only AeroBuildup has the bodies."""
    import aerosandbox as asb

    if de_deg:
        plane = plane.with_control_deflections({"elevator": de_deg})
    op = asb.OperatingPoint(atmosphere=atmosphere(cond["altitude_m"]), velocity=cond["V"], alpha=alpha_deg)
    if method == "buildup":
        r = asb.AeroBuildup(airplane=plane, op_point=op).run()
    elif method == "vlm":
        r = asb.VortexLatticeMethod(airplane=plane, op_point=op, **({"spanwise_resolution": 8, "chordwise_resolution": 6} | vlm)).run()
    elif method == "ll":
        r = asb.LiftingLine(airplane=plane, op_point=op, **({"spanwise_resolution": 8} | vlm)).run()
    elif method == "nll":
        r = asb.NonlinearLiftingLine(airplane=plane, op_point=op, **({"spanwise_resolution": 6} | vlm)).run()
    else:
        raise ValueError(method)
    return {"CL": num(r["CL"]), "CD": num(r["CD"]), "Cm": num(r["Cm"])}


def slopes(plane, cond: Dict[str, float], method: str = "buildup", a0_deg: float = 2.0, da_deg: float = 1.0,
           **vlm) -> Dict[str, float]:
    """Lift and moment slopes (per rad) at a0 from two angles, the static margin and the neutral point x.
    NP_x = x_ref + SM * c_ref, SM = -Cma / CLa (x is aft)."""
    lo, hi = (run_aero(plane, cond, a0_deg + s * da_deg, method, **vlm) for s in (-1, 1))
    d = math.radians(2 * da_deg)
    cla, cma = (hi["CL"] - lo["CL"]) / d, (hi["Cm"] - lo["Cm"]) / d
    cl0 = 0.5 * (hi["CL"] + lo["CL"]) - cla * math.radians(a0_deg)
    sm = -cma / cla
    return {"CLa": cla, "Cma": cma, "SM": sm, "NP_x": plane.xyz_ref[0] + sm * cond["c_ref"], "CL0": cl0,
            "alpha_L0_deg": -math.degrees(cl0 / cla)}


def trim_buildup(plane, cond: Dict[str, float], guess=(2.0, 0.0)) -> Dict[str, float]:
    """Level flight with AeroBuildup: alpha and elevator (TE down positive, deg) for L = W and Cm = 0 about plane.xyz_ref."""
    from scipy.optimize import fsolve

    def res(x):
        r = run_aero(plane, cond, x[0], "buildup", x[1])
        return [r["CL"] - cond["CL_req"], r["Cm"]]

    a, de = fsolve(res, guess)
    r = run_aero(plane, cond, a, "buildup", de)
    return _trim_row(r, a, de, cond)


def trim_flow5(lab: Lab, cond: Dict[str, float], cg_x: Optional[float] = None, guess=(2.0, 0.0)) -> Dict[str, float]:
    """Level flight with the flow5 DAVE-ML model (about the MRP). With cg_x the moment is moved to that x:
    Cm' = Cm + CL * (x' - x_mrp) / c_ref. Elevator in deg, TE down positive (JSBSim sign of the file)."""
    from scipy.optimize import fsolve

    x_mrp = lab.ref["xyz_ref"][0]
    shift = 0.0 if cg_x is None else (cg_x - x_mrp) / cond["c_ref"]

    def get(x):
        o = lab.flow5_eval(x[0], de=math.radians(x[1]))
        return o["CL"], o["CD"], o["Cm"] + o["CL"] * shift

    a, de = fsolve(lambda x: [get(x)[0] - cond["CL_req"], get(x)[2]], list(guess))
    cl, cd, cm = get([a, de])
    return _trim_row({"CL": cl, "CD": cd, "Cm": cm}, a, de, cond)


def _trim_row(r, alpha, de, cond):
    d = r["CD"] * cond["q"] * cond["s_ref"]
    return {"alpha_deg": float(alpha), "de_deg": float(de), "CL": r["CL"], "CD": r["CD"], "L_over_D": r["CL"] / r["CD"],
            "drag_N": d, "P_aero_W": d * cond["V"]}


def foil_polar(foil_name: str, coords, re: float, alphas_deg) -> Dict[str, np.ndarray]:
    """NeuralFoil section polar (CL, CD, CM) of foil coordinates."""
    import aerosandbox as asb

    af = asb.Airfoil(name=foil_name, coordinates=np.asarray(coords, float))
    a = np.asarray(alphas_deg, float)
    r = af.get_aero_from_neuralfoil(alpha=a, Re=re, mach=0.0, model_size="large")
    return {"alpha": a, "CL": np.asarray(r["CL"]), "CD": np.asarray(r["CD"]), "CM": np.asarray(r["CM"])}


# --------------------------------------------------------------------------------------------- the flow5 T7 reference
def t7_reference(lab: Lab) -> Dict[str, Any]:
    """The flow5 T7 trim point. Mass, CG and inertia: the mass model that the sweep flew (`analysis_mass_model` of the
    aero dataset, written by the flow5 spoke). Speed and angle of attack: the description of the aero dataset. After a
    sheet change the air vehicle mass model can differ from this one until the next flow5 sweep."""
    import re
    from aeroplan import adh_doc, assets

    out: Dict[str, Any] = {}
    air = adh_doc.node(lab.acft.doc, "1.2")["muas"]
    ref_id = next((d["ref"] for d in air.get("datasets", []) if d["role"] == "aero"), None)
    entry = next((e for e in assets.register(lab.acft.doc, create=False) if e["id"] == ref_id), {})
    mp = (entry.get("dataset") or {}).get("analysis_mass_model")
    if mp:
        out.update(mass_kg=mp["mass"]["value"], cg=mp["cg"]["xyz"], inertia=mp["inertia"])
    else:
        out["error"] = "the aero dataset holds no analysis_mass_model: run a new flow5 sweep and build"
    desc = ((lab.flow5_data or {}).get("file_header", {}).get("description") or {}).get("value", "")
    m = re.search(r"V (\d+\.?\d*) m/s, alpha (-?\d+\.?\d*) deg, CL (\d+\.?\d*)", desc)
    if m:
        out.update(V=float(m.group(1)), alpha_deg=float(m.group(2)), CL=float(m.group(3)))
    return out


# ------------------------------------------------------------------------------------------------------ plot style
#: one colour for each model, the same in every plot (categorical slots 1 to 3 of the validated palette; grey = a reference)
COLORS = {"flow5": "#2a78d6", "AeroBuildup": "#eb6834", "VLM": "#1baf7a", "Athena": "#52514e", "Sheet": "#8a8984"}


def style() -> None:
    """Matplotlib defaults: thin lines, a quiet grid, no top and right frame, light surface."""
    import matplotlib.pyplot as plt

    plt.rcParams.update({
        "figure.facecolor": "#fcfcfb", "axes.facecolor": "#fcfcfb", "savefig.facecolor": "#fcfcfb",
        "axes.edgecolor": "#c9c8c2", "axes.labelcolor": "#52514e", "xtick.color": "#52514e", "ytick.color": "#52514e",
        "text.color": "#0b0b0b", "axes.titlecolor": "#0b0b0b", "axes.spines.top": False, "axes.spines.right": False,
        "axes.grid": True, "grid.color": "#e6e5e0", "grid.linewidth": 0.8, "axes.axisbelow": True,
        "lines.linewidth": 2.0, "lines.markersize": 6, "legend.frameon": False, "figure.dpi": 110,
        "axes.titlesize": 11, "axes.titleweight": "bold", "axes.titlelocation": "left", "font.size": 10})


# --------------------------------------------------------------------------------------- planform and sheet numbers
def mac_geometry(lab: Lab, wing: str = "wing_main") -> Dict[str, float]:
    """Mean aerodynamic chord of a surface of the spec, the x of its leading edge and of its quarter chord (design frame)."""
    s = next(s for s in lab.spec["surfaces"] if s["name"] == wing)
    y = np.array([x["xyz_le"][1] for x in s["sections"]])
    z = np.array([x["xyz_le"][2] for x in s["sections"]])
    xle = np.array([x["xyz_le"][0] for x in s["sections"]])
    ch = np.array([x["chord"] for x in s["sections"]])
    t = np.concatenate([[0.0], np.cumsum(np.hypot(np.diff(y), np.diff(z)))])      # distance along the span
    tt = np.linspace(0.0, t[-1], 2001)
    c, x = np.interp(tt, t, ch), np.interp(tt, t, xle)
    area = np.trapezoid(c, tt)
    mac = np.trapezoid(c ** 2, tt) / area
    x_le = np.trapezoid(x * c, tt) / area
    return {"mac_m": float(mac), "x_le_m": float(x_le), "x_quarter_m": float(x_le + 0.25 * mac)}


def design_targets(root: Path = ROOT) -> Dict[str, Dict[str, Any]]:
    """Sheet 'Design Targets' of the Master Specification Sheet: {"athena": {label: value}, "reuben": {label: value}}.
    Text values such as '1.036 m^2' give the first number."""
    import re

    import openpyxl

    wb = openpyxl.load_workbook(root / "tools" / "specsheet" / "sheets" / "currawong" / "Currawong Master Specification Sheet.xlsx", data_only=True)
    ws = wb[next(n for n in wb.sheetnames if n.strip() == "Design Targets")]

    def number(v):
        if isinstance(v, (int, float)):
            return float(v)
        m = re.search(r"-?\d+\.?\d*(?:[eE]-?\d+)?", str(v)) if v is not None else None
        return float(m.group()) if m else None

    out: Dict[str, Dict[str, Any]] = {"athena": {}, "reuben": {}}
    for r in range(5, 50):
        for key, lc, vc in (("athena", 5, 6), ("reuben", 10, 11)):
            label, val = ws.cell(r, lc).value, ws.cell(r, vc).value
            if isinstance(label, str) and val is not None:
                out[key][label.strip()] = number(val) if key == "athena" or number(val) is not None else val
    return out
